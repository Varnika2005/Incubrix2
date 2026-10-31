import openpyxl

wb = openpyxl.load_workbook("IncuBrix_2027_Growth_Engineer_Option_3_AEO_SEO_GEO_Submission_Workbook_COMPLETED.xlsx", data_only=False)

print("Sheet summary:")
for name in wb.sheetnames:
    ws = wb[name]
    non_empty = sum(1 for r in ws.iter_rows(values_only=True) if any(v is not None for v in r))
    print(f"  {name:<20}: {non_empty} rows")

ws0 = wb['00_START']
print("\n00_START Key Values:")
print("  Candidate ID:", ws0['B5'].value)
print("  Full Name   :", ws0['B6'].value)
print("  Email       :", ws0['B7'].value)
print("  Degree      :", ws0['B8'].value)
print("  Repo URL    :", ws0['B12'].value)
print("  Videos URL  :", ws0['B14'].value)
print("  Declaration :", ws0['B21'].value)
print("  Checklist F5-F10:", [ws0[f'F{r}'].value for r in range(5, 11)])

ws4 = wb['04_SELF_CHECK']
print("\n04_SELF_CHECK Values:")
print("  C6..C12:", [ws4[f'C{r}'].value for r in range(6, 13)])
print("  C24..C28:", [ws4[f'C{r}'].value for r in range(24, 29)])

ws2 = wb['02_QUESTIONS']
print(f"\n02_QUESTIONS: {ws2.max_row} rows (Q-001..Q-100 populated)")
print("  First question:", ws2['B5'].value)
print("  Last question :", ws2['B104'].value)

ws2b = wb['02b_FIXES']
print(f"\n02b_FIXES: {ws2b.max_row} rows (ISS-001..ISS-022 populated)")
print("  First fix:", ws2b['A5'].value, "-", ws2b['C5'].value)
print("  Last fix :", ws2b['A26'].value, "-", ws2b['C26'].value)

ws2c = wb['02c_PAGES_CITATIONS']
print(f"\n02c_PAGES_CITATIONS: {ws2c.max_row} rows (10 Pages + 30 Citations populated)")
