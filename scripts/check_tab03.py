import openpyxl

wb = openpyxl.load_workbook('submission/workbook/IncuBrix_2027_Growth_Engineer_Option_3_AEO_SEO_GEO_Submission_Workbook.xlsx', data_only=True)
ws3 = wb['03_ENGINE_LAUNCH']
print("=== 03_ENGINE_LAUNCH Rows 6-31 ===")
for r in range(6, 32):
    row_vals = [f"Col{c}: {ws3.cell(row=r, column=c).value}" for c in range(1, 9) if ws3.cell(row=r, column=c).value is not None]
    if row_vals:
        print(f"Row {r}: " + " | ".join(row_vals))
